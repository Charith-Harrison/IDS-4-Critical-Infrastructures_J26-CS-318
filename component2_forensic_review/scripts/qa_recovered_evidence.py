"""Independent bounded QA of final candidate evidence. Reads original record offsets.

No rescanning/model fitting/source writes. Validates all candidate ADUs and available
paired response ADUs against immutable PCAP bytes, plus metadata and log-reference rules.
Frame ordinals are supplied from the streaming parser, not independently recounted here.
"""
from pathlib import Path
import csv,json,re,struct,hashlib,importlib.util,sys
from collections import Counter,defaultdict
from openpyxl import load_workbook
import numpy as np
import pandas as pd

sys.dont_write_bytecode=True
O=Path(__file__).resolve().parents[1];E=O/'evidence';RAW=Path(r'G:\modbus cic 2023\Modbus Dataset')
PARSER=Path(__file__).with_name('pcapio.py')  # preserved SHA256-identical copy of inventoried C3 parser
spec=importlib.util.spec_from_file_location('independent_candidate_qa_pcapio',PARSER)
pc=importlib.util.module_from_spec(spec);sys.modules[spec.name]=pc;spec.loader.exec_module(pc)

def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def integer(x):return int(float(x)) if '.' in str(x) else int(x)
errors=[];counts=Counter();opened={};headers={}
source=O/'recon_request_candidates.csv';before=sha(source)
rows=list(csv.DictReader(source.open(encoding='utf-8',newline='')))
logs={r['log_id']:r for r in csv.DictReader((E/'original_recon_log_rows.csv').open(encoding='utf-8',newline=''))}
def fail(row,kind,detail):errors.append({'candidate_id':row['candidate_id'],'kind':kind,'detail':detail})
def get_record(path,offset):
    if path not in opened:
        opened[path]=path.open('rb');headers[path]=pc.open_pcap(opened[path])
    fh=opened[path];h=headers[path];fh.seek(offset);hh=fh.read(16)
    sec,frac,caplen,origlen=struct.unpack(h.endian+'IIII',hh);data=fh.read(caplen)
    _,seg=pc.decode_l2l3l4(data,h.linktype)
    return sec*1_000_000+frac*1_000_000//h.ticks_per_second,caplen,seg

identities=[];tcpids=[];bygroup=defaultdict(list);refs_seen=set()
for row in rows:
    cid=row['candidate_id'];identities.append((row['source_relative_path'],row['frame_identity'],row['adu_offset_in_tcp_payload']))
    tcpids.append(tuple(row[k] for k in ['capture_run','src_ip','dst_ip','src_tcp_port','dst_tcp_port','tcp_connection_generation_observed','tcp_sequence','adu_hex']))
    bygroup[tuple(row[k] for k in ['capture_run','src_ip','dst_ip','unit_id'])].append(row)
    try:
        path=Path(row['source_reference']);rel=path.relative_to(RAW).as_posix()
        if rel!=row['source_relative_path']:fail(row,'request_path','absolute/relative mismatch')
        us,caplen,seg=get_record(path,integer(row['pcap_record_offset']))
        off=integer(row['adu_offset_in_tcp_payload']);adu=bytes.fromhex(row['adu_hex'])
        if us!=integer(row['timestamp_us']):fail(row,'request_timestamp','PCAP header differs')
        if caplen!=integer(row['packet_size']):fail(row,'request_size','captured frame length differs')
        if seg.payload[off:off+len(adu)]!=adu:fail(row,'request_bytes','ADU differs from source bytes')
        if (seg.src_ip,seg.dst_ip,seg.src_port,seg.dst_port)!=(row['src_ip'],row['dst_ip'],integer(row['src_tcp_port']),integer(row['dst_tcp_port'])):fail(row,'request_endpoints','source endpoints differ')
        tid,proto,length=struct.unpack('!HHH',adu[:6]);unit=adu[6];fc=adu[7]
        if proto!=0 or len(adu)!=12 or length!=6 or fc not in [1,2,3,4]:fail(row,'read_request_shape','not canonical FC1..4 12-byte request ADU')
        vals=(tid,unit,fc,length,int.from_bytes(adu[8:10],'big'),int.from_bytes(adu[10:12],'big'))
        expected=tuple(integer(row[k]) for k in ['transaction_id','unit_id','wire_function_code','modbus_len','address','quantity_value'])
        if vals!=expected:fail(row,'request_fields','parsed source fields differ')
        if seg.seq+off!=integer(row['tcp_sequence']):fail(row,'request_tcp_sequence','sequence+offset differs')
        counts['request_adus_source_checked']+=1
        if row.get('response_frame_identity',''):
            rp=RAW/row['response_source_reference']
            rus,rsz,rseg=get_record(rp,integer(row['response_pcap_record_offset']))
            roff=integer(row['response_adu_offset_in_tcp_payload']);radu=bytes.fromhex(row['response_adu_hex'])
            if rseg.payload[roff:roff+len(radu)]!=radu:fail(row,'response_bytes','ADU differs from source bytes')
            if rus!=integer(row['response_timestamp_us']):fail(row,'response_timestamp','PCAP header differs')
            if (rseg.src_ip,rseg.dst_ip,rseg.src_port,rseg.dst_port)!=(row['dst_ip'],row['src_ip'],integer(row['dst_tcp_port']),integer(row['src_tcp_port'])):fail(row,'response_endpoints','not reversed request endpoints')
            if int.from_bytes(radu[:2],'big')!=tid or radu[6]!=unit or radu[7]&127!=fc:fail(row,'response_pair_fields','TID/unit/basefunction differs')
            if rus-us!=integer(row['response_delta_us']) or not 0<=rus-us<=5_000_000:fail(row,'response_delta','incorrect/noncausal/out-of-bound candidate pair')
            counts['response_adus_source_checked']+=1
    except Exception as ex:fail(row,'source_read_or_parse',repr(ex))
    # Preserve XLSX ;row= locators: only semicolons followed by source path separate refs.
    refs=re.split(r';(?=attack/)',row['attack_log_reference'])
    if len(refs)!=integer(row['log_match_count']):fail(row,'log_reference_count',f'{len(refs)} parsed refs vs declared {row["log_match_count"]}')
    for ref in refs:
        refs_seen.add(ref);log=logs.get(ref)
        if log is None:fail(row,'log_reference','unknown original row '+ref);continue
        if log['target_ip'] and log['target_ip']!=row['dst_ip']:fail(row,'log_target','target not request destination')
        if log['range_value'] and integer(log['range_value'])!=integer(row['quantity_value']):fail(row,'log_range','range not request quantity')
        if not abs(integer(log['timestamp_candidate_utc_us'])-integer(row['timestamp_us']))<=25_000:fail(row,'log_delta','outside declared contextual bound')
        counts['original_log_references_checked']+=1
    unique=integer(row['log_match_count'])==1
    if (row['label_state']=='supported_contextual')!=unique:fail(row,'label_state','uniqueness does not match contextual/ambiguous state')
    if row['evidence_supported_recon_request'].lower()!=str(unique).lower():fail(row,'evidence_flag','ambiguous row promoted or supported row suppressed')
    if row['exact_clock_contract_verified'].lower()!='false':fail(row,'clock_claim','unproven exact clock marked verified')
    n=integer(row['previous_context_n'])
    if not 0<=n<=20:fail(row,'history_size','outside 0..20')
    if n==0 and row['previous_request_delta_s']:fail(row,'history_missingness','gap exists with no previous context')
    if n and (not row['previous_request_delta_s'] or float(row['previous_request_delta_s'])<0):fail(row,'history_chronology','missing/negative prior-request gap')
for group,rr in bygroup.items():
    times=[integer(x['timestamp_us']) for x in rr]
    pos=[integer(x['sequence_position_after_ordering']) for x in rr]
    if times!=sorted(times):errors.append({'group':group,'kind':'candidate_order','detail':'timestamps decrease'})
    if pos!=list(range(len(rr))):errors.append({'group':group,'kind':'candidate_positions','detail':'not exact zero-based count after ordering'})
# Independently read referenced rows from original log files, not merely the derived index.
needed=defaultdict(dict)
for ref in refs_seen:
    log=logs[ref];needed[log['source_path']][ref]=log
def verify_original_log(ref,expected,actual):
    mapping={'timestamp_original':'Timestamp','target_ip':'TargetIP','attack':'Attack','transaction_id_original':'TransactionID'}
    for left,right in mapping.items():
        value=str(actual.get(right,''))
        if value!=expected[left]:errors.append({'log_reference':ref,'kind':'original_log_row_field','detail':left+': '+repr(value)+' != '+repr(expected[left])})
    counts['original_source_log_rows_checked']+=1
for relative,refs in needed.items():
    path=RAW/relative
    if path.suffix.lower()=='.csv':
        byrow={integer(v['source_row']):(k,v) for k,v in refs.items()}
        with path.open(encoding='utf-8-sig',errors='replace',newline='') as fh:
            for no,actual in enumerate(csv.DictReader(fh),2):
                if no in byrow:
                    ref,log=byrow[no];verify_original_log(ref,log,actual)
    elif path.suffix.lower()=='.xlsx':
        wb=load_workbook(path,read_only=True,data_only=True)
        bysheet=defaultdict(dict)
        for ref,log in refs.items():
            sheet=ref.split('#sheet=',1)[1].rsplit(';row=',1)[0]
            bysheet[sheet][integer(log['source_row'])]=(ref,log)
        for sheet,byrow in bysheet.items():
            ws=wb[sheet];it=ws.iter_rows(values_only=True);header=[str(x) for x in next(it)]
            for no,values in enumerate(it,2):
                if no in byrow:
                    ref,log=byrow[no];verify_original_log(ref,log,dict(zip(header,values)))
        wb.close()
for path,fh in opened.items():fh.close()
result={'candidate_rows':len(rows),'candidate_ids_unique':len({r['candidate_id'] for r in rows})==len(rows),'duplicate_raw_adu_identities':len(identities)-len(set(identities)),'duplicate_same_tcp_generation_sequence_payload_identities':len(tcpids)-len(set(tcpids)),'counts':dict(counts),'logical_candidate_streams':len(bygroup),'distinct_original_log_row_references':len(refs_seen),'label_states':dict(Counter(r['label_state'] for r in rows)),'errors':errors,'all_performed_checks_pass':len(errors)==0,'candidate_file_sha256_before':before,'candidate_file_sha256_after':sha(source),'limitations':['Frame ordinals were not independently recounted; exact record-offset bytes were independently read.','TCP generation and uniqueness of response attribution are not proven by endpoint/TID matching.','The global logger clock contract remains unverified; contextual labels are not exact ground truth.','Causal history construction is reviewed in code; this bounded check does not rescan all preceding requests.']}
(E/'quality_raw_evidence_checks.json').write_text(json.dumps(result,indent=2),encoding='utf-8')
print(json.dumps(result,indent=2),flush=True)
