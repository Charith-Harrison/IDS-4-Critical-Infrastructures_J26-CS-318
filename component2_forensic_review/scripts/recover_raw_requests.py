"""Independent source-byte audit; bounded single-segment Modbus ADU recovery.

Selects one canonical NIC/bridge view for every scenario with recon log evidence.
All original logs are scanned. Reassembled/missing ADUs are not guessed. This is
forensic candidate recovery, never final ML preprocessing or label propagation.
"""
import sys
sys.dont_write_bytecode=True
from pathlib import Path
import csv,json,struct,bisect,re,time,hashlib,importlib.util
from collections import Counter,defaultdict,deque
from datetime import datetime,timezone
import numpy as np
import pandas as pd
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parents[1]
E=OUT/'evidence'
RAW=Path(r'G:\modbus cic 2023\Modbus Dataset')
PCAPIO=Path(__file__).with_name('pcapio.py')  # preserved SHA256-identical copy of inventoried C3 parser
spec=importlib.util.spec_from_file_location('c2_readonly_pcapio',PCAPIO)
pcapio=importlib.util.module_from_spec(spec);sys.modules[spec.name]=pcapio;spec.loader.exec_module(pcapio)


def dump(name,value):
    (OUT/name).write_text(json.dumps(value,indent=2,default=str)+'\n',encoding='utf8')


def family(path):
    return 'external_February_recording' if '/external/' in path else ('compromised_IED_March_recording' if '/compromised-ied/' in path else 'compromised_SCADA_March_recording')


def parse_ts(s):
    try:
        return int(datetime.fromisoformat(str(s)).replace(tzinfo=timezone.utc).timestamp()*1000000)
    except (ValueError,TypeError):
        return None


def load_logs():
    recon=[]; other=defaultdict(list); counts=Counter();files=[]
    for path in sorted(RAW.rglob('*.csv')):
        rel=path.relative_to(RAW).as_posix();n=0
        with path.open(newline='',encoding='utf-8-sig',errors='replace') as fh:
            for rowno,r in enumerate(csv.DictReader(fh),2):
                n+=1;label=r.get('Attack','');ts=parse_ts(r.get('Timestamp',''));target=r.get('TargetIP','')
                counts['csv_rows']+=1
                if 'recon' in label.lower():
                    rr={'log_id':f'{rel}#row={rowno}','source_path':rel,'source_row':rowno,
                        'campaign':family(rel),'timestamp_original':r.get('Timestamp',''),
                        'timestamp_candidate_utc_us':ts,'target_ip':target,'attack':label,
                        'transaction_id_original':r.get('TransactionID',''),
                        'is_completion':'complete' in label.lower(),'source_format':'csv'}
                    recon.append(rr)
                elif ts is not None and ('brute force' in label.lower() or 'specific coil' in label.lower()):
                    # Literal row identity retained; no completion-as-initiator substitution.
                    other[(family(rel),target)].append((ts,rel,rowno,label,r.get('TransactionID','')))
        files.append({'path':rel,'rows':n})
    for path in sorted(RAW.rglob('*.xlsx')):
        rel=path.relative_to(RAW).as_posix();wb=load_workbook(path,read_only=True,data_only=True)
        for ws in wb:
            rows=ws.iter_rows(values_only=True);hdr=[str(x) for x in next(rows)]
            for rowno,values in enumerate(rows,2):
                r=dict(zip(hdr,values));label=str(r.get('Attack',''))
                counts['xlsx_rows']+=1
                if 'recon' in label.lower():
                    recon.append({'log_id':f'{rel}#sheet={ws.title};row={rowno}','source_path':rel,
                        'source_row':rowno,'campaign':family(rel),'timestamp_original':str(r.get('Timestamp','')),
                        'timestamp_candidate_utc_us':parse_ts(r.get('Timestamp')),
                        'target_ip':str(r.get('TargetIP','')),'attack':label,
                        'transaction_id_original':str(r.get('TransactionID','')),
                        'is_completion':'complete' in label.lower(),'source_format':'xlsx'})
        wb.close()
    for r in recon:
        try:t=int(r['transaction_id_original']);r['wire_tid_valid']=0<=t<=65535
        except ValueError:r['wire_tid_valid']=False
        r['range_value']=int(m.group(1)) if (m:=re.search(r'Range:\s*(\d+)',r['attack'])) else None
        r['clock_assumption']='naive logger time treated as UTC candidate; corroborated locally by packet proximity, global clock contract unproven'
    pd.DataFrame(recon).to_csv(E/'original_recon_log_rows.csv',index=False)
    # Exact duplicate signature references retained separately to avoid inflating event counts.
    starts=defaultdict(list)
    seen={}
    for r in recon:
        if not r['is_completion'] and r['timestamp_candidate_utc_us'] is not None:
            key=(r['campaign'],r['timestamp_candidate_utc_us'],r['target_ip'],r['attack'],r['transaction_id_original'])
            if key in seen:
                seen[key]['duplicate_source_log_ids'].append(r['log_id'])
            else:
                rr=dict(r);rr['duplicate_source_log_ids']=[];seen[key]=rr
                starts[r['campaign']].append(rr)
    for v in starts.values():v.sort(key=lambda x:x['timestamp_candidate_utc_us'])
    for k,v in other.items():v.sort()
    dump('evidence/original_log_inventory.json',{'files':files,'counts':counts,'recon_rows':len(recon),
        'distinct_start_signatures':len(seen),'other_brute_specific_rows':sum(map(len,other.values()))})
    return recon,starts,other


def nearest_log(starts,times,camp,us,dst,qty):
    rows=starts.get(camp,[]);ts=times.get(camp,[])
    lo=bisect.bisect_left(ts,us-25000);hi=bisect.bisect_right(ts,us+25000)
    candidates=[r for r in rows[lo:hi] if (not r['target_ip'] or r['target_ip']==dst)
                and (r['range_value'] is None or r['range_value']==qty)]
    return candidates


def qstats(v):
    a=np.array(v,float)
    return {'n':len(a),**({str(q):float(np.quantile(a,q)) for q in [0,.25,.5,.75,.95,1]} if len(a) else {})}


def main():
    t0=time.monotonic();recon,starts,other=load_logs()
    supplement='--hmi-prefix' in sys.argv
    times={k:[r['timestamp_candidate_utc_us'] for r in v] for k,v in starts.items()}
    other_times={k:[r[0] for r in v] for k,v in other.items()}
    p=pd.read_csv(E/'recon_labeled_packets_ALL.csv');p['processed_row_id']=np.arange(len(p))
    full_date=p.timestamp.astype(str).str.match(r'^\d{4}-\d{2}-\d{2}[T ]')
    pt=pd.to_datetime(p.timestamp.where(full_date),format='mixed',utc=True,errors='coerce')
    p['ts_us']=pt.dt.as_unit('ns').astype('int64')//1000
    pos=p[p.label.eq(1)].copy()
    positive_keys=defaultdict(list)
    for r in pos.itertuples():
        positive_keys[(r.file_source,r.ts_us,r.src_ip,r.dst_ip,r.func_code,r.modbus_len)].append(r.processed_row_id)
    # Single canonical view per scenario. SCADA bridge includes ALL rotations.
    paths=sorted((RAW/'attack/compromised-scada/substation-wide-capture').glob('*.pcap'),key=lambda x:int(x.stem.split('-')[-1]))
    paths=list((RAW/'attack/external/external-attacker/external-attacker-network-capture').glob('*.pcap'))+paths
    # Extra mirror views solely reconcile old positives; never become independent candidates.
    extra=[x for x in RAW.rglob('*.pcap') if x.name in set(pos.file_source) and x not in paths]
    allpaths=paths+extra
    if supplement:
        paths=[RAW/'attack/compromised-scada/scada-hmi/scada-hmi-network-capture/veth291ab9a-0.pcap']
        extra=[];allpaths=paths
    candidates=[];normal=[];rawpositive=[];examples=[];file_reports=[]
    last_request={};read_hist=defaultdict(lambda:deque(maxlen=20));tcp_seen=set();generation=defaultdict(int)
    prefix_metrics=Counter();candidate_payloads=set();window_counts=Counter();samplecount=Counter()
    columns=None
    for path in allpaths:
        canonical=path in paths;rel=path.relative_to(RAW).as_posix();camp=family(rel)
        counts=Counter();min_us=None;max_us=None
        with path.open('rb') as fh:
            h=pcapio.open_pcap(fh)
            for frame,sec,frac,caplen,origlen,data,offset in pcapio.iter_records(fh,h):
                counts['frames']+=1
                if data is None:counts['truncated_record']+=1;continue
                us=sec*1000000+frac*1000000//h.ticks_per_second
                min_us=us if min_us is None else min(min_us,us);max_us=us if max_us is None else max(max_us,us)
                if supplement and us>=1678870408842227:
                    counts['post_primary_bridge_start_frames_omitted_to_avoid_mirrors']+=1;continue
                kind,seg=pcapio.decode_l2l3l4(data,h.linktype)
                if seg is None or (seg.src_port!=502 and seg.dst_port!=502):continue
                counts['modbus_tcp_segments']+=1
                stream=(camp,seg.src_ip,seg.src_port,seg.dst_ip,seg.dst_port)
                if seg.flags&2 and not seg.flags&16:generation[stream]+=1
                if not seg.payload:continue
                raw=seg.payload;off=0
                while off+7<=len(raw):
                    tid,proto,length=struct.unpack_from('!HHH',raw,off)
                    if proto!=0 or length<2 or length>254:
                        counts['noncanonical_or_unframed_payloads']+=1;break
                    end=off+6+length
                    if end>len(raw):counts['incomplete_single_segment_adu']+=1;break
                    adu=raw[off:end];unit=adu[6];rawfc=adu[7];fc=rawfc&127
                    isreq=seg.dst_port==502;address=None;quantity=None;value=None
                    if isreq and fc in (1,2,3,4,5,6) and len(adu)>=12:
                        address=int.from_bytes(adu[8:10],'big')
                        if fc in (1,2,3,4):quantity=int.from_bytes(adu[10:12],'big')
                        else:value=int.from_bytes(adu[10:12],'big')
                    counts['complete_segment_adus']+=1
                    dedup=(stream,generation[stream],seg.seq+off,adu)
                    retrans=dedup in tcp_seen
                    if len(tcp_seen)>2000000:tcp_seen.clear();counts['dedup_cache_reset']+=1
                    tcp_seen.add(dedup)
                    rr={'source_reference':str(path),'source_relative_path':rel,'capture_run':camp,
                        'capture_chunk':path.name,'frame_identity':frame,'pcap_record_offset':offset,
                        'adu_offset_in_tcp_payload':off,'timestamp_us':us,
                        'timestamp':datetime.fromtimestamp(us/1e6,timezone.utc).isoformat(timespec='microseconds'),
                        'src_ip':seg.src_ip,'dst_ip':seg.dst_ip,'src_tcp_port':seg.src_port,'dst_tcp_port':seg.dst_port,
                        'tcp_sequence':seg.seq+off,'tcp_connection_generation_observed':generation[stream],
                        'tcp_generation_status':'observed SYN count; may start mid-connection',
                        'unit_id':unit,'func_code':fc,'wire_function_code':rawfc,'modbus_len':length,
                        'address':address,'quantity_value':quantity if quantity is not None else value,
                        'packet_size':caplen,'transaction_id':tid,'adu_hex':adu.hex(),
                        'is_request':isreq,'is_exception':bool(rawfc&128),
                        'exception_code':adu[8] if rawfc&128 and len(adu)>8 else None,
                        'raw_source_verified':True,'canonical_view':canonical,'retransmission_candidate':retrans,
                        'extraction_status':'complete_adu_within_single_captured_tcp_segment'}
                    ids=positive_keys.get((path.name,us,seg.src_ip,seg.dst_ip,fc,length),[])
                    if ids:
                        for rid in ids:
                            rp=dict(rr);rp['original_processed_row_id']=rid
                            pair=last_request.get((camp,seg.dst_ip,seg.dst_port,seg.src_ip,seg.src_port,tid,unit,fc))
                            if pair and not 0<=us-pair['timestamp_us']<=5_000_000:
                                counts['stale_or_negative_old_positive_pairs_rejected']+=1;pair=None
                            rp['matching_request_frame']=pair['frame_identity'] if pair else None
                            rp['matching_request_source']=pair['source_relative_path'] if pair else ''
                            rp['matching_request_adu_hex']=pair['adu_hex'] if pair else ''
                            rp['matching_request_address']=pair['address'] if pair else None
                            rp['matching_request_quantity_value']=pair['quantity_value'] if pair else None
                            rp['matching_request_timestamp_us']=pair['timestamp_us'] if pair else None
                            rp['request_response_delta_us']=us-pair['timestamp_us'] if pair else None
                            rp['pairing_state']='same_reversed_TCP_endpoints_TID_unit_function' if pair else 'unmatched'
                            okey=(camp,pair['dst_ip'] if pair else seg.src_ip)
                            ots=other_times.get(okey,[]);i=bisect.bisect_left(ots,pair['timestamp_us'] if pair else us)
                            opts=[j for j in [i-1,i] if 0<=j<len(ots)]
                            near=min(opts,key=lambda j:abs(ots[j]-(pair['timestamp_us'] if pair else us))) if opts else None
                            if near is not None:
                                row=other[okey][near]
                                rp['nearest_nonrecon_log_reference']=f'{row[1]}#row={row[2]}'
                                rp['nearest_nonrecon_log_attack']=row[3]
                                rp['nearest_nonrecon_log_delta_us']=(pair['timestamp_us'] if pair else us)-row[0]
                            rawpositive.append(rp)
                    if isreq:
                        last_request[(camp,seg.src_ip,seg.src_port,seg.dst_ip,seg.dst_port,tid,unit,fc)]=rr
                    elif canonical:
                        pair=last_request.get((camp,seg.dst_ip,seg.dst_port,seg.src_ip,seg.src_port,tid,unit,fc))
                        if pair and pair.get('candidate_index') is not None:
                            target=candidates[pair['candidate_index']]
                            if 0<=us-pair['timestamp_us']<=5_000_000:
                                target.update({'response_source_reference':rel,'response_frame_identity':frame,
                                    'response_timestamp_us':us,'response_adu_hex':adu.hex(),
                                    'response_wire_function_code':rawfc,'response_exception_code':rr['exception_code'],
                                    'response_delta_us':us-pair['timestamp_us'],
                                    'response_pairing':'reversed_TCP_endpoints_same_TID_unit_base_function'})
                    if canonical and isreq and fc in (1,2,3,4) and quantity is not None and not retrans:
                        histkey=(camp,seg.src_ip,seg.dst_ip,unit)
                        hist=read_hist[histkey]
                        if hist and (us-hist[-1][0]>60_000_000 or us<hist[-1][0]):hist.clear()
                        addresses=[x[1] for x in hist];dts=np.diff([x[0] for x in hist])/1e6 if len(hist)>1 else []
                        rr.update({'previous_context_n':len(hist),
                            'previous_request_delta_s':(us-hist[-1][0])/1e6 if hist else None,
                            'address_delta':address-hist[-1][1] if hist else None,
                            'absolute_address_delta':abs(address-hist[-1][1]) if hist else None,
                            'unique_addresses_previous20':len(set(addresses)),
                            'register_span_previous20':max(addresses)-min(addresses) if addresses else None,
                            'monotonic_increasing_fraction_previous20':float(np.mean(np.diff(addresses)>0)) if len(addresses)>1 else None,
                            'function_transitions_previous20':sum(x[2]!=y[2] for x,y in zip(hist,list(hist)[1:])),
                            'previous20_interval_mean_s':float(np.mean(dts)) if len(dts) else None,
                            'previous20_interval_cv':float(np.std(dts)/np.mean(dts)) if len(dts) and np.mean(dts)>0 else None,
                            'previous20_requests_per_s':(len(hist)-1)/((hist[-1][0]-hist[0][0])/1e6) if len(hist)>1 and hist[-1][0]>hist[0][0] else None,
                            'causal_context_max_idle_gap_s':60})
                        matches=nearest_log(starts,times,camp,us,seg.dst_ip,quantity)
                        if matches:
                            rr['attack_log_reference']=';'.join(x['log_id'] for x in matches)
                            rr['log_match_count']=len(matches)
                            rr['log_timestamp_original']=';'.join(sorted(set(x['timestamp_original'] for x in matches)))
                            rr['log_target_ip']=';'.join(sorted(set(x['target_ip'] for x in matches)))
                            rr['log_transaction_id_original']=';'.join(sorted(set(x['transaction_id_original'] for x in matches)))
                            rr['log_delta_us_min']=min(us-x['timestamp_candidate_utc_us'] for x in matches)
                            rr['log_delta_us_max']=max(us-x['timestamp_candidate_utc_us'] for x in matches)
                            exact=[x for x in matches if x['wire_tid_valid'] and int(x['transaction_id_original'])==tid and x['target_ip']==seg.dst_ip]
                            rr['label_state']='supported_contextual'
                            rr['wire_tid_matching_log_start_count']=len(exact)
                            rr['exact_clock_contract_verified']=False
                            rr['matching_confidence']='target/range/time within25ms; '+('literal wire TID agrees' if exact else 'TID absent/nonmatching/repeated')
                            rr['evidence_supported_recon_request']=True
                            rr['candidate_id']=f'raw_recon_{len(candidates):06d}'
                            rr['candidate_index']=len(candidates)
                            candidates.append(rr)
                        else:
                            # Same shape controls from the actual raw stream. This proves
                            # absence of nearby recon evidence, never blanket benign truth.
                            shape=(camp,fc,quantity,length,caplen)
                            samplecount[shape]+=1
                            if samplecount[shape]<=40 or samplecount[shape]%500==0:
                                rr['control_state']='no_recon_log_start_matching_target_range_within25ms'
                                rr['benign_ground_truth_verified']=False
                                normal.append(rr)
                        hist.append((us,address,fc))
                    off=end
                if off<len(raw) and off+7>len(raw):counts['trailing_short_payload_bytes']+=len(raw)-off
        file_reports.append({'source':rel,'canonical':canonical,'size':path.stat().st_size,
            'min_timestamp_us':min_us,'max_timestamp_us':max_us,'counts':dict(counts)})
        print(json.dumps({'file':rel,'canonical':canonical,'frames':counts['frames'],
            'recon_candidates_so_far':len(candidates),'raw_positive_matches_so_far':len(rawpositive),
            'seconds':round(time.monotonic()-t0,1)}),flush=True)
    c=pd.DataFrame(candidates)
    if len(c):
        c=c.sort_values(['capture_run','src_ip','dst_ip','unit_id','timestamp_us','source_relative_path','frame_identity','adu_offset_in_tcp_payload'])
        c['sequence_position_after_ordering']=c.groupby(['capture_run','src_ip','dst_ip','unit_id']).cumcount()
        c.drop(columns=['candidate_index'],errors='ignore').to_csv(OUT/('evidence/hmi_prefix_recon_request_candidates.csv' if supplement else 'recon_request_candidates.csv'),index=False)
    pd.DataFrame(normal).to_csv(E/('hmi_prefix_normal_request_context.csv' if supplement else 'normal_request_context.csv'),index=False)
    rp=pd.DataFrame(rawpositive)
    if not supplement:rp.to_csv(E/'raw_positive_packet_matches.csv',index=False)
    metrics={'canonical_sources':len(paths),'extra_views_for_old_label_reconciliation':len(extra),
        'original_log_rows_recon':len(recon),'distinct_recon_start_signatures':sum(map(len,starts.values())),
        'candidate_requests':len(c),'label_states':c.label_state.value_counts().to_dict() if len(c) else {},
        'candidate_campaigns':c.capture_run.value_counts().to_dict() if len(c) else {},
        'candidate_function_codes':c.func_code.value_counts().to_dict() if len(c) else {},
        'candidate_quantities':c.quantity_value.value_counts().to_dict() if len(c) else {},
        'candidate_addresses':c.address.value_counts().to_dict() if len(c) else {},
        'candidates_with_response':int(c.get('response_frame_identity',pd.Series(dtype=float)).notna().sum()),
        'raw_positive_matches':len(rp),'unique_original_positive_rows_matched':rp.original_processed_row_id.nunique() if len(rp) else 0,
        'raw_positive_wire_function_distribution':rp.wire_function_code.value_counts().to_dict() if len(rp) else {},
        'raw_positive_exception_code_distribution':rp.exception_code.value_counts().to_dict() if len(rp) else {},
        'raw_positive_nearest_nonrecon_log_labels':rp.get('nearest_nonrecon_log_attack',pd.Series(dtype=str)).value_counts().to_dict(),
        'controls_saved':len(normal),'file_reports':file_reports,
        'limitations':['Single-segment complete ADUs only; incomplete/unframed payload counts retained, no full TCP reassembly',
            '25ms contextual log join is sensitivity-bounded evidence, not proof of global logger clock synchronization',
            'No independent-run claim from capture rotations or repeated tool invocations',
            'Controls from attack recording are unlabeled local lookalikes, not automatically benign'],
        'wall_seconds':time.monotonic()-t0}
    dump('evidence/hmi_prefix_raw_recovery_metrics.json' if supplement else 'evidence/raw_recovery_metrics.json',metrics)
    print(json.dumps({k:v for k,v in metrics.items() if k!='file_reports'},indent=2,default=str))


if __name__=='__main__':main()
