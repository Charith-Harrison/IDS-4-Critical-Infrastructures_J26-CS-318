"""Stream the supplied .xlsx workbooks into CSV working copies (read-only on originals).

The supplied files are named *.csv.xlsx but are genuine OOXML workbooks (zip magic PK).
They are opened read-only and never modified. Every value is written as text exactly as
openpyxl returns it, so nanosecond timestamp strings are preserved verbatim.
"""
from __future__ import annotations

import csv
import datetime as dt
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from common import load_config, resolve, sha256_file, write_json  # noqa: E402


def to_text(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.isoformat()
    if isinstance(v, float):
        return repr(v)
    return str(v)


def convert(src: Path, dst: Path) -> dict:
    t0 = time.time()
    wb = load_workbook(src, read_only=True, data_only=True)
    sheets = [ws.title for ws in wb.worksheets]
    ws = wb.worksheets[0]
    n = 0
    dst.parent.mkdir(parents=True, exist_ok=True)
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            w.writerow([to_text(v) for v in row])
            n += 1
    wb.close()
    return {"source": str(src), "sheets": sheets, "rows_including_header": n,
            "output": str(dst), "seconds": round(time.time() - t0, 1)}


def main():
    cfg = load_config()
    raw_dir = resolve(cfg["paths"]["original_dataset_dir"])
    out_dir = resolve(cfg["paths"]["raw_copy_dir"])
    manifest = {"original_dataset_dir": str(raw_dir), "files": []}
    for src in sorted(raw_dir.glob("*.xlsx")):
        dst = out_dir / (src.name.replace(".csv.xlsx", "").replace(".xlsx", "") + ".csv")
        info = {"original_sha256_before": sha256_file(src), "original_size": src.stat().st_size,
                "original_mtime": src.stat().st_mtime}
        if dst.exists() and "--force" not in sys.argv:
            info.update({"source": str(src), "output": str(dst), "skipped_existing": True})
        else:
            print("converting", src.name, flush=True)
            info.update(convert(src, dst))
        info["original_sha256_after"] = sha256_file(src)
        assert info["original_sha256_before"] == info["original_sha256_after"], "original modified!"
        info["copy_sha256"] = sha256_file(dst)
        manifest["files"].append(info)
        print(info, flush=True)
    write_json(manifest, out_dir / "raw_manifest.json")


if __name__ == "__main__":
    main()
