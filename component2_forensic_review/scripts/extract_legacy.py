"""Read-only original XLSX extraction. Outputs are forensic derivatives, not raw data."""
import csv, datetime, hashlib, json, time
from pathlib import Path
from openpyxl import load_workbook

OUT = Path(__file__).resolve().parents[1] / 'evidence'
SRC = Path(r'D:\VSCode\GitHub\4th Year Cyber\C2')
def sha(p):
    h=hashlib.sha256()
    with p.open('rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()
def text(v):
    if v is None:return ''
    if isinstance(v,datetime.datetime):return v.isoformat()
    if isinstance(v,float):return repr(v)
    return str(v)
manifest=[]
for name in ['attack_events_mapped.xlsx','recon_labeled_packets_ALL.csv.xlsx','recon_sequences_ALL.csv.xlsx']:
    src=SRC/name; dst=OUT/(name.replace('.csv.xlsx','.csv').replace('.xlsx','.csv'))
    t=time.time(); before=sha(src)
    wb=load_workbook(src,read_only=True,data_only=True)
    info={'source':str(src),'source_sha256_before':before,'sheets':[s.title for s in wb.worksheets],'output':str(dst),'rows':0}
    with dst.open('w',encoding='utf-8',newline='') as f:
        w=csv.writer(f)
        for row in wb.worksheets[0].iter_rows(values_only=True):
            w.writerow([text(v) for v in row]); info['rows']+=1
    wb.close();info.update(source_sha256_after=sha(src),output_sha256=sha(dst),seconds=time.time()-t)
    assert before==info['source_sha256_after']
    manifest.append(info)
    print(json.dumps(info),flush=True)
(OUT/'legacy_extraction_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
